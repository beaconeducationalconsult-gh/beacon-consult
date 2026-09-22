from pathlib import Path
import json,zipfile,re,hashlib
from collections import Counter
from docx import Document
from PIL import Image
ROOT=Path(__file__).resolve().parent.parent;OUT=ROOT/'files';DOWNLOADS=ROOT/'downloads';DOWNLOADS.mkdir(exist_ok=True)
records=json.loads((ROOT/'source/assessment_records.json').read_text())
weekly=[r for r in records if 'week' in r];tests=[r for r in records if 'test' in r]
assert len(weekly)==24 and len(tests)==4
assert {(r['grade'],r['week']) for r in weekly}=={(g,w) for g in (4,5) for w in range(1,13)}
assert len(list(OUT.rglob('*.docx')))==9 and len(list(OUT.rglob('*.png')))==24
balance=Counter()
for r in records:
    assert (len(r['mcq']),len(r['essays']))==((20,10) if 'week' in r else (15,15))
    assert all(len(e)==3 for e in r['essays'])
    assert all(sum(t['marks'] for t in e)==(4 if 'week' in r else 6) for e in r['essays'])
    for q in r['mcq']:
        assert len(q['options'])==4 and len(set(q['options']))==4
        assert q['options'].count(q['answer'])==1
        assert q['options']['ABCD'.index(q['letter'])]==q['answer']
        balance[q['letter']]+=1
    if 'test' in r:
        allowed=6 if r['test']=='Mid_Term' else 12
        assert {q['source_week'] for q in r['mcq']}==set(range(1,allowed+1))
        assert all(1<=t['source_week']<=allowed for e in r['essays'] for t in e)
        assert len({(q['q'],tuple(sorted(q['options']))) for q in r['mcq']})==15
    else:
        assert (OUT/r['map']).exists()
        assert len({q['q'] for q in r['mcq']})==20
embedded=0
for p in OUT.rglob('*.docx'):
    with zipfile.ZipFile(p) as z:
        assert z.testzip() is None
        embedded+=len([n for n in z.namelist() if n.startswith('word/media/')])
    d=Document(p);text='\n'.join(p.text for p in d.paragraphs)
    assert text.strip() and 'TODO' not in text
    if 'Weekly_Lessons' in p.name:
        assert len(d.inline_shapes)==12
        assert sum(p.text=='Section A — Multiple-choice questions' for p in d.paragraphs)==12
        assert sum(p.text=='Section B — 10 structured essay questions' for p in d.paragraphs)==12
        assert len([p for p in d.paragraphs if re.match(r'^\([abc]\) ',p.text)])==360
        assert 'MCQ answers and explanations' not in text
    if p.name.endswith('_Test.docx'):
        assert len([p for p in d.paragraphs if re.match(r'^\d+\. ',p.text)])==30
        assert len([p for p in d.paragraphs if re.match(r'^\([abc]\) ',p.text)])==45
        assert 'MCQ answers and explanations' not in text
assert embedded==24
for p in OUT.rglob('*.png'):
    im=Image.open(p);assert im.size in ((2200,1800),(1264,843),(1536,1024));im.verify()
report={'batch':2,'scope':'Science B4 and B5, Weeks 1–12 and both tests; overall original request still incomplete','word_documents':9,'weekly_units':24,'new_weekly_units_since_batch_1':22,'week_1_units_reincluded':2,'standalone_mindmaps':24,'embedded_mindmaps':embedded,'weekly_mcqs':480,'weekly_three_part_questions':240,'test_papers':4,'test_mcqs':60,'test_three_part_questions':60,'total_mcqs':540,'total_three_part_questions':300,'total_structured_subparts':900,'answer_position_distribution':dict(balance),'checks_passed':['DOCX ZIP integrity','Required weekly and test counts','Four distinct MCQ alternatives and one keyed answer','MCQ key letters match options','No duplicated MCQ stems within a weekly bank','No duplicate MCQ items within a test paper','Test source weeks within declared scope','Correct 30-question test numbering and 45 subparts','Expected structured-question marks','Separate teacher keys','24 embedded maps and 24 valid PNGs'],'review_limits':['No full Word pagination render','No independent teacher or scientific sign-off','B4 mapping interpretations require teacher confirmation','Tests draw on weekly banks and are not secure unseen papers','All 12 revised B4 illustrated maps visually inspected; B5 maps remain original text-based maps']}
manifest=json.loads((ROOT/'source/illustrated_maps.json').read_text())
book=Document(OUT/'B4/Science/B4_Science_Weekly_Lessons_W01-W12.docx')
for item,shape in zip(manifest['weeks'],book.inline_shapes):
    data=(OUT/item['path']).read_bytes()
    assert hashlib.sha256(data).hexdigest()==item['sha256']
    rid=shape._inline.graphic.graphicData.pic.blipFill.blip.embed
    assert book.part.related_parts[rid].blob==data
report['illustrated_maps']={'B4_Science':12,'B5_Science':0}
report['checks_passed'].append('All 12 embedded B4 illustrations match standalone PNG hashes')
(OUT/'QUALITY_CHECKS.json').write_text(json.dumps(report,indent=2,ensure_ascii=False),encoding='utf-8')
(OUT/'SOURCE_CROSSWALK.json').write_text((ROOT/'source/source_crosswalk.json').read_text(),encoding='utf-8')
(OUT/'FILE_MANIFEST.txt').write_text('\n'.join(f'{p.relative_to(OUT)} | {p.stat().st_size:,} bytes' for p in sorted(OUT.rglob('*')) if p.is_file() and p.name!='FILE_MANIFEST.txt')+'\n',encoding='utf-8')
archive=DOWNLOADS/'Beacon_Term1_Batch2_Science_B4_B5.zip'
with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
    for p in sorted(OUT.rglob('*')):
        if p.is_file():z.write(p,arcname=str(p.relative_to(OUT)))
with zipfile.ZipFile(archive) as z:assert z.testzip() is None
info={'filename':archive.name,'bytes':archive.stat().st_size,'sha256':hashlib.sha256(archive.read_bytes()).hexdigest()}
(DOWNLOADS/'archive-info.json').write_text(json.dumps(info,indent=2))
print(json.dumps(report,indent=2));print(json.dumps(info,indent=2))

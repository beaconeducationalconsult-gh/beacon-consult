"""Embed reviewed standalone B4 maps without rebuilding questions or answer keys."""
from pathlib import Path
import hashlib,json
from docx import Document
from PIL import Image
ROOT=Path(__file__).resolve().parent.parent
folder=ROOT/'files/B4/Science'
book=folder/'B4_Science_Weekly_Lessons_W01-W12.docx'
doc=Document(book)
assert len(doc.inline_shapes)==12
before=[p.text for p in doc.paragraphs]
manifest=[]
for week,shape in enumerate(doc.inline_shapes,1):
    path=folder/'Mind-maps'/f'B4_Science_Week_{week:02d}.png'
    image=Image.open(path);image.verify()
    w,h=Image.open(path).size
    rid=shape._inline.graphic.graphicData.pic.blipFill.blip.embed
    doc.part.related_parts[rid]._blob=path.read_bytes()
    shape.height=round(shape.width*h/w)
    manifest.append({'week':week,'path':str(path.relative_to(ROOT/'files')),'size':[w,h],'sha256':hashlib.sha256(path.read_bytes()).hexdigest()})
doc.save(book)
assert [p.text for p in Document(book).paragraphs]==before
(ROOT/'source/illustrated_maps.json').write_text(json.dumps({'grade':4,'subject':'Science','weeks':manifest,'status':'AI-generated illustrations with editorial corrections; teacher-review draft','provenance':'Week 1 approved preview; Weeks 2–11 generated for their topics; Week 12 composed from approved common-bean artwork.','review':'All twelve visually inspected; corrected insect leg count, separation diagram, hub label, plant captions and germination layout. No independent scientific sign-off.'},indent=2)+'\n')
print('Embedded 12 B4 illustrations; lesson/assessment paragraph text unchanged.')

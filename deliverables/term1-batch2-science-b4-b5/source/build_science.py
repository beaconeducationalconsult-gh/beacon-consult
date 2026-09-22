"""Produce complete Science B4 and B5 weekly books and assessments."""
from pathlib import Path
import sys,json,random,copy
ROOT=Path(__file__).resolve().parent.parent
sys.path.insert(0,str(ROOT.parent/'term1-batch1'/'source'))
import build as layout
sys.path.insert(0,str(Path(__file__).resolve().parent))
from content import LESSONS
from docx.shared import Inches
OUT=ROOT/'files';OUT.mkdir(exist_ok=True)
records=[]

def doc(title,subtitle,student=False):
    d=layout.newdoc(title,subtitle,student)
    d.sections[0].header.paragraphs[0].text='BEACON EDUCATIONAL CONSULT / 2026–2027 / TERM 1 / SCIENCE BATCH 2'
    return d
para=layout.para;bullet=layout.bullet;save=layout.save

def mcq_bank(g,w):
    m=LESSONS[(g,w)];qs=[]
    for j,field in enumerate(['q1','q2']):
        for i,c in enumerate(m['cards']):
            t=dict(q=c[field],answer=c['answer'],wrong=c['wrong'],method=c['reason'],source_week=w,concept=c['answer'])
            qs.append(layout.shuffled(t,g*10000+w*100+j*10+i))
    return qs

def weekly_essays(m):
    return [[dict(q=f'Explain the term “{c["answer"]}”.',answer=c['fact'],marks=1),dict(q=c['q2'],answer=c['answer'],marks=1),dict(q=c['extend'],answer=c['reason'],marks=2)] for c in m['cards']]

def put_week_essays(d,es):
    d.add_heading('Section B — 10 structured essay questions',2)
    para(d,'Answer all parts (a), (b), (c). Write in a separate exercise book. Suggested marks per question: (a) 1, (b) 1, (c) 2, total 4. Use accurate terms and explain relationships or evidence where asked.')
    for i,parts in enumerate(es,1):
        d.add_heading(f'Question {i}',3)
        for label,t in zip('abc',parts):para(d,f'({label}) {t["q"]}')

def put_key(d,qs,es,start=1,test=False):
    layout.mcq_key(d,qs);d.add_heading('Structured-response marking guide',2)
    if test:para(d,'Each part is worth 2 marks. Part (a): 1 for accurate meaning and 1 for a supporting scientific point. Part (b): 1 for the correct identification and 1 for a reason. Part (c): 2 for two relevant points or a clear relationship supported by evidence/example. Accept scientifically correct equivalents; do not require exact wording.')
    else:para(d,'Part (a): 1 for accurate meaning. Part (b): 1 for correct identification. Part (c): 2 for two relevant points or a clear relationship supported by evidence/example. Accept equivalent accurate wording and examples. The indicative notes may contain more detail than is needed for full marks.')
    for i,parts in enumerate(es,start):
        d.add_heading(f'Question {i}',3)
        for label,t in zip('abc',parts):para(d,f'({label}) [{t["marks"]} mark(s)] {t["answer"]}')

def source_table(d,g):
    d.add_heading('Weekly sequence and source crosswalk',1)
    table=d.add_table(rows=1,cols=3);table.style='Light Shading Accent 1'
    for cell,text in zip(table.rows[0].cells,['Week / topic','Source reference → used reference','Editorial decision']):cell.text=text
    for w in range(1,13):
        m=LESSONS[(g,w)];cells=table.add_row().cells
        cells[0].text=f'{w}. {m["title"]}';cells[1].text=f'{m["source_codes"]} → {m["codes"]}';cells[2].text=m['mapping_note']

def build_grade(g):
    folder=OUT/f'B{g}'/'Science'
    d=doc(f'Basic {g} Science — Term 1','Weeks 1–12 • Weekly lesson notes, study notes, visual mind-maps and practice',True)
    para(d,'BATCH 2 — COMPLETE WEEKLY SEQUENCE FOR THIS GRADE, TEACHER-REVIEW DRAFT. Week 1 is included from Batch 1 so this Science book is self-contained. This is not a claim of official curriculum approval or independent teacher certification.')
    para(d,'Each teaching week contains one consolidated lesson unit, suggested teaching sessions, study notes, one visual mind-map, 20 four-option MCQs and 10 three-part structured questions. These are weekly units, not separate full daily lesson plans. Keep the separate teacher answer book away from pupil copies.')
    para(d,'Source sequence: first-term-library/BASIC '+str(g)+' SCHEME OF LEARNING.pdf, Science table page '+str(24 if g==4 else 26)+'. Weeks 13–15 are reserved in this scheme for revision/examinations/vacation, so no new weekly content is added there. Mid-term assumes Weeks 1–6; end-of-term assumes Weeks 1–12. Confirm the school’s calendar.')
    para(d,'Important: the B4 source contains several indicator mismatches. The crosswalk states the interpretation instead of silently claiming exact alignment. In particular, B4 Week 5 is an introductory materials-properties lesson inferred from its Materials block, not a verified B4 indicator. Obtain teacher confirmation before use. Supporting vocabulary sometimes extends beyond the wording of the indicator.')
    source_table(d,g)
    k=doc(f'Basic {g} Science — Teacher answers','Weeks 1–12, mid-term and end-of-term • Separate marking guide')
    para(k,'Use with the matching Batch 2 student books and papers. Assess scientific meaning rather than memorised wording. Review source interpretations, question difficulty, safety and test timing before use. Tests draw on the weekly learning banks and are not advertised as secure unseen examinations.')
    for w in range(1,13):
        m=LESSONS[(g,w)];qs=mcq_bank(g,w);es=weekly_essays(m)
        path=folder/'Mind-maps'/f'B{g}_Science_Week_{w:02d}.png'
        if g == 4 and (ROOT/'source/illustrated_maps.json').exists():
            assert path.exists(), f'Missing reviewed illustration: {path}'
        else:
            layout.mindmap(path,m['title'],g,'Science',w,[(c['answer'],c['fact']+'.') for c in m['cards']],f'Sequence: BASIC {g} Science p{m["page"]}. See source crosswalk for code interpretations. Teacher-review draft.')
        d.add_page_break();d.add_heading(f'Week {w} — {m["title"]}',1)
        para(d,f'Scheme reference: {m["source_codes"]}. Reference used: {m["codes"]}. {m["mapping_note"]}')
        d.add_heading('Learning outcomes',2)
        for t in [f'Explain the main ideas in {m["title"].lower()} using accurate vocabulary.','Use observations or models to identify examples and justify explanations.','Apply the ideas to a familiar situation, recognise a misconception and state relevant safety precautions.']:bullet(d,t)
        d.add_heading('Teacher lesson notes',2)
        para(d,'Suggested weekly structure: three 40-minute teaching sessions, plus distributed practice and feedback in remaining timetable periods. Adjust to actual lesson periods. The full question bank should not be imposed as one 40-minute task.')
        para(d,'Prior-knowledge check: ask learners for a familiar example related to the week. Record what they think happens and why. Retain one initial claim to revisit after the activity.')
        sessions=[('Session 1 — Observe and define','5 min: retrieve related knowledge; 10 min: show a model, picture or safe specimen; 15 min: pairs record observations and discuss the first five terms; 10 min: check explanations with short oral questions.'),('Session 2 — Investigate and explain','5 min: recall vocabulary; 10 min: model a justified explanation; 15 min: carry out the specific activity below, recording real observations; 10 min: introduce the remaining terms and compare evidence.'),('Session 3 — Apply and assess','5 min: revisit an initial misconception; 10 min: explain the map connections; 15 min: attempt selected application questions and structured parts; 10 min: review explanations and plan reteaching or extension.')]
        for h,t in sessions:d.add_heading(h,3);para(d,t)
        d.add_heading('Topic-specific activity and resources',2);para(d,m['activity'])
        para(d,'Resources: the teacher-selected examples, pictures or apparatus described above; board, pencils, exercise books and the supplied map. If a practical cannot be completed, use a clearly labelled photograph or model and do not report predicted results as observed results.')
        d.add_heading('Safety',2);para(d,m['safety'])
        d.add_heading('Support, extension and feedback',2)
        para(d,'Support: provide picture labels and a word bank, read unfamiliar words aloud, and let learners explain orally before writing. Extension: compare two related concepts, suggest a counterexample or critique a model. Feedback: ask why a distractor is wrong and use the answer to identify the misconception.')
        d.add_heading('Study notes',1)
        for c in m['cards']:d.add_heading(c['answer'],2);para(d,c['fact']+'.');para(d,c['reason'])
        para(d,'Independent study: reconstruct the map from memory, give a fresh example for three branches and answer one explanation question without copying the notes. Correct errors after discussing the reason with a teacher or partner.')
        d.add_page_break();d.add_heading(f'Week {w} — Visual mind-map',1);d.add_picture(str(path),width=Inches(6.9));para(d,'Read each branch, explain the link to the central topic, and add your own example. The separate high-resolution PNG supports projection or larger printing. Concept-map connections do not imply every process occurs in a single linear order.')
        d.add_page_break();d.add_heading(f'Week {w} — Practice and assessment',1);layout.put_mcq(d,qs);put_week_essays(d,es)
        k.add_page_break();k.add_heading(f'Week {w} — Answers',1);put_key(k,qs,es)
        records.append(dict(grade=g,subject='Science',week=w,title=m['title'],mcq=qs,essays=es,map=str(path.relative_to(OUT))))
    save(d,folder/f'B{g}_Science_Weekly_Lessons_W01-W12.docx')
    for label,maxweek in [('Mid_Term',6),('End_of_Term',12)]:
        exam=doc(f'Basic {g} Science — {label.replace("_"," ")} Test',f'Term 1 • Weeks 1–{maxweek} • 30 numbered questions',True)
        para(exam,'Answer ALL questions. Section A: 15 MCQs × 1 mark = 15. Section B: 15 structured essay questions, each with (a), (b), (c), 2 marks per part = 90. Total = 105 marks. Percentage = score ÷ 105 × 100. Use a separate answer booklet.')
        para(exam,'Suggested administration: two supervised sessions, approximately 45 minutes for Section A and 135 minutes for Section B, adjusted by the school. The requested 45 structured subparts make a long paper. Confirm timing, reading support and accommodations before use.')
        qs=[];es=[];blueprint=[]
        for i in range(15):
            w=i%maxweek+1;m=LESSONS[(g,w)];cards=m['cards'];index=(i*3+(2 if maxweek==12 else 0))%10;c=cards[index]
            t=dict(q=c['q2'],answer=c['answer'],wrong=c['wrong'],method=c['reason'],source_week=w,concept=c['answer'])
            qs.append(layout.shuffled(t,g*100000+maxweek*1000+i))
            ca=cards[(index+2)%10];cb=cards[(index+5)%10];cc=cards[(index+8)%10]
            parts=[dict(q=f'Explain the meaning of “{ca["answer"]}”, including a supporting scientific point.',answer=ca['fact']+'. '+ca['reason'],marks=2,source_week=w),dict(q=cb['q2']+' Give a scientific reason for your answer.',answer=cb['answer']+'. '+cb['reason'],marks=2,source_week=w),dict(q=cc['extend'],answer=cc['reason'],marks=2,source_week=w)]
            es.append(parts);blueprint.append(f'{i+1}/{i+16} → Week {w}: {m["title"]}')
        layout.put_mcq(exam,qs);exam.add_heading('Section B — Structured essay questions',2)
        para(exam,'Answer all parts. Every part carries 2 marks. Give explanations, evidence or examples as requested; short accurate paragraphs are sufficient.')
        for i,parts in enumerate(es,16):
            exam.add_heading(f'{i}. Explain and apply the following related scientific ideas.',3)
            for l,t in zip('abc',parts):para(exam,f'({l}) {t["q"]}')
        save(exam,folder/f'B{g}_Science_{label}_Test.docx')
        k.add_page_break();k.add_heading(label.replace('_',' ')+' — Test answers',1)
        para(k,'Blueprint (MCQ / structured question → week). Items draw on the learning banks; review or replace them if an unseen assessment is required.')
        for row in blueprint:para(k,row)
        put_key(k,qs,es,16,True)
        records.append(dict(grade=g,subject='Science',test=label,mcq=qs,essays=es,blueprint=blueprint))
    save(k,folder/f'B{g}_Science_Teacher_Answers.docx')

def guide():
    d=doc('START HERE — Science B4 and B5','Batch 2 • Complete Weeks 1–12 for these two subject–grade combinations')
    d.add_heading('Included in this archive',1)
    for t in ['B4 Science: Weeks 1–12, mid-term test, end-of-term test and separate teacher answers.','B5 Science: Weeks 1–12, mid-term test, end-of-term test and separate teacher answers.','Each weekly unit: lesson notes, study notes, one visual mind-map, 20 four-option MCQs and 10 structured essay questions with parts (a), (b), (c).','Each test: 15 MCQs plus 15 three-part structured questions, numbered 1–30.','Nine editable Word documents, including this guide, and 24 standalone high-resolution mind-maps.']:bullet(d,t)
    para(d,'Week 1 for both grades is included from Batch 1. This batch therefore adds 22 new weekly units and four new tests; it does not claim 24 entirely new units.')
    d.add_heading('Cumulative status after Batch 2',1)
    table=d.add_table(rows=1,cols=3);table.style='Light Shading Accent 1'
    for cell,t in zip(table.rows[0].cells,['Subject','Grades','Delivered status']):cell.text=t
    for row in [('Mathematics','B4–B8','Weeks 1–12 and both tests in Batch 1'),('Science','B4–B5','Weeks 1–12 and both tests in Batch 2'),('Science','B6–B8','Week 1 only in Batch 1; later weeks and both tests outstanding'),('Computing','B4–B8','Week 1 only in Batch 1; later weeks and both tests outstanding')]:
        for cell,t in zip(table.add_row().cells,row):cell.text=t
    para(d,'The full original request is still incomplete. No automatic background generation of the outstanding material is implied by this delivery.')
    d.add_heading('Sources and interpretation',1)
    para(d,'Weekly order comes from the supplied BASIC 4 scheme Science page 24 and BASIC 5 page 26, supported by data/curriculum/science_B4_curriculum_db_clean.json and science_B5_curriculum_db_clean.json. References are local source references, not a new official curriculum certification.')
    para(d,'B4 Weeks 3–6 and 8–10 contain source codes that are absent or inconsistent in the local curriculum descriptions. The weekly crosswalk explicitly shows replacements or interpretations. B4 Week 5 is an editorial introductory properties unit within the Materials block, supported by B5 property vocabulary rather than a verified B4 indicator. Confirm this choice with the school before teaching or formal assessment. B5 references match the local descriptions. Some supporting terms extend beyond the short indicator wording.')
    para(d,'B4 Week 5 and B5 Week 3 share core material-property content, with a more comparative B5 activity. Some related concepts recur intentionally. These are consolidated weekly units with suggested sessions, not complete individual daily lesson plans.')
    d.add_heading('Assessment and printing',1)
    para(d,'Mid-term covers Weeks 1–6; end-of-term covers Weeks 1–12. Each paper totals 105 marks: 15 MCQ marks plus 90 structured-response marks. Convert score to a percentage using score/105 × 100. The large question count needs long or split administration. Tests use selected application questions and recombined structured parts from the weekly learning banks; they are not claimed to be unseen or secure. Adapt where school policy requires fresh items.')
    para(d,'Open and edit in Word or a compatible editor. Keep Teacher_Answers separate from student copies. Print selected weeks rather than the entire book if preferred. Pupils answer in separate exercise books or test booklets. PNG maps are also embedded in the books and may be projected or printed larger.')
    d.add_heading('Safety and review',1)
    para(d,'Practical activities are low-risk or model-based. Never drink classroom-treated water, look directly at the Sun, taste experimental seeds, handle unknown chemicals, or collect hazardous waste. Heating, disinfection, gases and fire equipment are adult-controlled or represented only by diagrams. Follow school safety guidance. Recorded results must be actual observations, not invented values.')
    para(d,'Automated checks verify DOCX package integrity, counts, four distinct options and one keyed answer, map embedding, and test numbering. They do not independently validate every scientific statement, teaching suitability, accessibility or Word pagination. See QUALITY_CHECKS.json for the exact check results. All documents remain teacher-review drafts.')
    if (ROOT/'source/illustrated_maps.json').exists():
        d.add_heading('Illustrated-map revision — B4 Science',1)
        para(d,'B4 Science Weeks 1–12 have illustrated maps embedded in Word and supplied as PNGs. B5 maps remain text-based pending revision. Batch 1 is unchanged. AI-generated artwork with editorial corrections requires teacher review.')
    save(d,OUT/'START_HERE_Batch_2_Science_B4_B5.docx')
    (OUT/'README.txt').write_text('BATCH 2 — SCIENCE B4 AND B5\n\nIncludes Weeks 1–12, both tests, separate answers and 24 mind-maps. Week 1 is repeated from Batch 1 for self-contained books.\n\nNine Word documents. Open START_HERE_Batch_2_Science_B4_B5.docx.\n\nOutstanding overall: Science B6–B8 later weeks and both tests; Computing B4–B8 later weeks and both tests.\n\nTeacher-review drafts. Read source crosswalk and safety notes before use.\n',encoding='utf-8')

    if (ROOT/'source/illustrated_maps.json').exists():
        with (OUT/'README.txt').open('a',encoding='utf-8') as f:
            f.write('\nILLUSTRATED REVISION\n'+'B4 Science Weeks 1–12 have illustrated maps embedded in Word and supplied as PNGs. B5 maps remain text-based pending revision. Batch 1 is unchanged. AI-generated artwork with editorial corrections requires teacher review.'+'\n')

if __name__=='__main__':
    for g in (4,5):build_grade(g);print(f'Science B{g} complete: 12 weeks, 2 tests, answers.',flush=True)
    guide();(ROOT/'source'/'assessment_records.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
    (ROOT/'source'/'source_crosswalk.json').write_text(json.dumps([{k:v for k,v in LESSONS[(g,w)].items() if k in ['grade','week','title','source_codes','codes','mapping_note','page']} for g in (4,5) for w in range(1,13)],ensure_ascii=False,indent=2),encoding='utf-8')

"""Build the partial pack in a durable repository-backed directory."""
from pathlib import Path
import json,random,math
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from math_content import MODULES,WEEKS,PAGES,task
from starter_content import DATA
ROOT=Path(__file__).resolve().parent.parent
OUT=ROOT/'files';OUT.mkdir(exist_ok=True)
FONT='/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
BOLD='/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'
records=[]
def wrap(draw,text,font,width):
    lines=[];line=''
    for word in text.split():
        test=(line+' '+word).strip()
        if draw.textlength(test,font=font)>width and line:lines.append(line);line=word
        else:line=test
    if line:lines.append(line)
    return lines

def mindmap(path,title,g,subject,w,nodes,source):
    im=Image.new('RGB',(2200,1800),'#F3F7FB');d=ImageDraw.Draw(im)
    d.rounded_rectangle((20,20,2180,1778),radius=35,outline='#BACBD8',width=4)
    d.text((60,48),f'B{g} / {subject.upper()} / TERM 1 / WEEK {w}',font=ImageFont.truetype(BOLD,45),fill='#17324D')
    rows=math.ceil(len(nodes)/2);gap=28;top=170;height=int((1660-top-(rows-1)*gap)/rows)
    colors=['#008785','#3975B5','#805EA8','#B67617','#B24762'];boxes=[]
    for i,node in enumerate(nodes):
        side=i%2;row=i//2;x=60 if side==0 else 1480;y=top+row*(height+gap);color=colors[row%5];cx=790 if side==0 else 1410;nx=x+660 if side==0 else x
        d.line([(cx,900),(cx-35 if side==0 else cx+35,900),(nx+30 if side==0 else nx-30,y+height//2),(nx,y+height//2)],fill=color,width=7)
        boxes.append((x,y,color,node))
    d.rounded_rectangle((790,705,1410,1080),radius=38,fill='#17324D')
    size=38
    while True:
        font=ImageFont.truetype(BOLD,size);lines=wrap(d,title,font,550)
        if len(lines)*(size+10)<=260:break
        size-=1
    for i,line in enumerate(lines):d.text((1100,745+i*(size+10)),line,font=font,anchor='mt',fill='white')
    d.text((1100,1030),'CONNECT • EXPLAIN • APPLY',font=ImageFont.truetype(FONT,24),anchor='mt',fill='#89E2D7')
    for x,y,color,(heading,body) in boxes:
        d.rounded_rectangle((x,y,x+660,y+height),radius=22,fill='white',outline=color,width=4)
        hf=ImageFont.truetype(BOLD,30);yy=y+18
        for line in wrap(d,heading,hf,620):d.text((x+20,yy),line,font=hf,fill=color);yy+=37
        fs=27
        while True:
            f=ImageFont.truetype(FONT,fs);lines=wrap(d,body,f,620)
            if yy+len(lines)*(fs+8)<y+height-12:break
            fs-=1
            if fs<18:raise ValueError(('Diagram overflow',title,heading))
        for line in lines:d.text((x+20,yy),line,font=f,fill='#243C50');yy+=fs+8
    d.text((60,1705),'READ A BRANCH → EXPLAIN IT → GIVE AN EXAMPLE → CONNECT IT TO ANOTHER BRANCH',font=ImageFont.truetype(FONT,24),fill='#17324D')
    d.text((60,1750),source,font=ImageFont.truetype(FONT,18),fill='#536B7B')
    path.parent.mkdir(parents=True,exist_ok=True);im.save(path,optimize=True)

def newdoc(title,subtitle,student=False):
    d=Document();s=d.sections[0];s.page_width=Inches(8.27);s.page_height=Inches(11.69)
    s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(.65)
    normal=d.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(5)
    for name,size,color in [('Title',28,'17324D'),('Heading 1',20,'17324D'),('Heading 2',15,'007F83'),('Heading 3',12,'17324D')]:
        st=d.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string(color)
    s.header.paragraphs[0].text='BEACON EDUCATIONAL CONSULT / 2026–2027 / TERM 1 / REBUILT EDITION'
    footer=s.footer.paragraphs[0];footer.alignment=2;footer.add_run('Emergency teacher-review draft | ')
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
    d.add_heading(title,0);d.add_paragraph(subtitle,'Subtitle')
    if student:d.add_paragraph('Name: __________________________________  Class: __________\nDate: __________________  Teacher: ________________________')
    return d

def para(d,text):d.add_paragraph(text)
def bullet(d,text):d.add_paragraph(text,'List Bullet')
def save(d,p):p.parent.mkdir(parents=True,exist_ok=True);d.save(p)
def shuffled(t,seed):
    options=[t['answer']]+t['wrong'];random.Random(seed).shuffle(options)
    return {**t,'options':options,'letter':'ABCD'[options.index(t['answer'])]}

def math_mcq(skills,g,seed,n=20):
    r=random.Random(seed);out=[];seen=set();attempts=0
    while len(out)<n:
        skill=skills[len(out)%len(skills)];t=task(skill,g,r,attempts//len(skills));attempts+=1
        fp=(t['q'],tuple(sorted([t['answer']]+t['wrong'])))
        if fp in seen:
            assert attempts<5000,('Not enough distinct tasks',g,skills)
            continue
        seen.add(fp);out.append(shuffled(t,seed+attempts))
    return out

def math_essays(skills,g,seed,n=10):
    r=random.Random(seed)
    return [[task(skills[(i+j)%len(skills)],g,r,i+j) for j in range(3)] for i in range(n)]

def put_mcq(d,qs):
    d.add_heading('Section A — Multiple-choice questions',2)
    para(d,'Choose ONE correct answer for each question. Each question has four alternatives.')
    for i,t in enumerate(qs,1):
        para(d,f'{i}. {t["q"]}')
        para(d,'    '.join(f'{k}. {v}' for k,v in zip('ABCD',t['options'])))

def put_essays(d,essays,start=1):
    d.add_heading('Section B — Structured essay questions',2)
    para(d,'Answer every part (a), (b), (c). Show your method or explain your choice. Write in a separate answer booklet. Each part: 1 mark for valid working/explanation and 1 for the correct result; 6 marks per question.')
    for i,parts in enumerate(essays,start):
        d.add_heading(f'{i}. Apply the related skills and show working or give a reason for each answer.',3)
        for label,t in zip('abc',parts):para(d,f'({label}) {t["q"]}')

def mcq_key(d,qs):
    d.add_heading('MCQ answers and explanations',2)
    for i,t in enumerate(qs,1):para(d,f'{i}. {t["letter"]} — {t["answer"]}. {t["method"]}')

def essay_key(d,essays,start=1):
    d.add_heading('Structured-response marking guide',2)
    para(d,'Each part: 1 mark for a valid method/explanation and 1 mark for the correct result. Accept equivalent exact forms unless a specified form is required. Accept alternative valid methods. Apply the school’s agreed follow-through policy to arithmetic slips.')
    for i,parts in enumerate(essays,start):
        d.add_heading(f'Question {i}',3)
        for label,t in zip('abc',parts):para(d,f'({label}) {t["answer"]}. {t["method"]}')

SESSIONS=[('1 — Explore and diagnose','5 minutes: recall an earlier skill; 10: demonstrate a concrete or pictorial model; 15: pairs investigate and explain; 10: identify a misconception and set an exit task.'),('2 — Model and practise','5 minutes: review the exit task; 10: work an example aloud; 15: pairs solve parallel examples; 10: compare methods and correct errors.'),('3 — Connect representations','5 minutes: recall vocabulary; 10: explain the mind-map; 15: connect a model, a drawing and a symbolic calculation; 10: independent practice.'),('4 — Apply and explain','5 minutes: estimate or predict; 10: solve a contextual example; 15: attempt selected structured-question parts; 10: compare reasoning and check units.'),('5 — Check and respond','5 minutes: retrieve the main rule; 15: use selected MCQs diagnostically; 10: explain distractors and misconceptions; 10: reteach or extend in response to evidence.')]

def build_math(g):
    folder=OUT/f'B{g}'/'Mathematics'
    d=newdoc(f'Basic {g} Mathematics','Term 1 • Weeks 1–12 • Lesson notes, study notes, mind-maps and weekly practice',True)
    para(d,'EMERGENCY TEACHER-REVIEW DRAFT — REBUILT EDITION. Topic order follows the supplied Alpha Examinations first-term scheme. Notes, examples and diagrams are newly prepared. This is not an official curriculum approval or a guarantee of full indicator-by-indicator validation.')
    para(d,'Each week contains one consolidated weekly lesson unit, one visual mind-map, 20 four-option MCQs and 10 structured questions with parts (a), (b), (c). Daily sessions are suggested, not separate full daily lesson plans. Answers are in the separate teacher book.')
    para(d,'The source reserves Weeks 13–15 for revision/examinations/vacation. Mid-term assessment assumes Weeks 1–6; end-of-term assumes Weeks 1–12. Confirm school dates and pacing. Some task structures repeat with new values for fluency practice.')
    d.add_heading('Weekly contents',1)
    for w,skills in enumerate(WEEKS[g],1):para(d,f'{w:02d} — '+' / '.join(MODULES[k]['title'] for k in skills))
    k=newdoc(f'Basic {g} Mathematics — Teacher answers','Weekly practice, mid-term test and end-of-term test • Keep separate from pupil copies')
    para(k,'Review difficulty, topic coverage and wording before formal assessment. These are worked draft keys, not independent teacher certification.')
    for w,skills in enumerate(WEEKS[g],1):
        title=' / '.join(MODULES[s]['title'] for s in skills)
        nodes=[]
        if len(skills)==1:
            m=MODULES[skills[0]];nodes=[(f'Key idea {i+1}',s) for i,s in enumerate(m['notes'])]
            example=task(skills[0],g,random.Random(g*100+w),0)
            nodes += [('Worked connection',example['method']),('Avoid this mistake',m['error'])]
        else:
            for s in skills:
                m=MODULES[s];nodes += [(m['title'],m['notes'][0]+' '+m['notes'][1]),('Apply: '+m['title'],m['notes'][-1])]
        path=folder/'Mind-maps'/f'B{g}_Maths_Week_{w:02d}.png'
        mindmap(path,title,g,'Mathematics',w,nodes,f'Sequence: BASIC {g} scheme p{PAGES[g]}. Newly prepared explanation and diagram; teacher-review draft.')
        qs=math_mcq(skills,g,100000*g+w*100);essays=math_essays(skills,g,200000*g+w*100)
        d.add_page_break();d.add_heading(f'Week {w} — {title}',1)
        para(d,f'Sequence source: BASIC {g} SCHEME OF LEARNING.pdf, first-term Mathematics, page {PAGES[g]}. One consolidated weekly unit.')
        d.add_heading('Learning outcomes',2)
        for s in skills:bullet(d,'Explain and apply '+MODULES[s]['title'].lower()+', solve examples and justify the method.')
        d.add_heading('Teaching notes and suggested sessions',2)
        para(d,'Suggested schedule: five 40-minute sessions, adjusted to the school timetable. Spread the question bank through the week; it is not intended to be completed in one lesson.')
        for h,t in SESSIONS:d.add_heading(h,3);para(d,t)
        d.add_heading('Topic-specific activities',2)
        for s in skills:bullet(d,MODULES[s]['activity'])
        para(d,'Resources: board, pencils, number cards, counters, paper strips and squared paper as appropriate. Prior-knowledge check: ask a learner to explain a representation used in the first example before introducing a new rule.')
        para(d,'Support: reduce number size without changing the concept; offer a partially completed chart or model; allow oral reasoning before writing. Extension: create a counterexample to a false rule and explain why it fails.')
        d.add_heading('Study notes and worked examples',2)
        for i,s in enumerate(skills):
            m=MODULES[s];d.add_heading(m['title'],3)
            for t in m['notes']:bullet(d,t)
            example=task(s,g,random.Random(g*10000+w*100+i),i)
            para(d,'Worked example: '+example['q']);para(d,'Solution: '+example['method']);para(d,'Common misconception: '+m['error'])
        para(d,'Independent study: read the map, cover it and reconstruct the branches. Solve an unused question from each skill and explain the reasoning to a partner. Consult the teacher key only after an honest attempt.')
        d.add_page_break();d.add_heading(f'Week {w} — Visual mind-map',1);d.add_picture(str(path),width=Inches(6.9));para(d,'The separate high-resolution PNG can be projected or printed larger. Explain each branch and connect it with an example.')
        d.add_page_break();d.add_heading(f'Week {w} — Practice',1);put_mcq(d,qs);put_essays(d,essays)
        k.add_page_break();k.add_heading(f'Week {w} — Answers',1);mcq_key(k,qs);essay_key(k,essays)
        records.append(dict(grade=g,subject='Mathematics',week=w,mcq=qs,essays=essays,map=str(path.relative_to(OUT))))
    save(d,folder/f'B{g}_Mathematics_Weekly_Lessons_W01-W12.docx')
    for label,maxweek in [('Mid_Term',6),('End_of_Term',12)]:
        exam=newdoc(f'Basic {g} Mathematics — {label.replace("_"," ")} Test',f'Term 1 • Weeks 1–{maxweek} • 30 numbered questions',True)
        para(exam,'Answer ALL questions. Section A: 15 MCQs × 1 mark = 15 marks. Section B: 15 questions × 3 parts × 2 marks = 90 marks. Total: 105 marks. Percentage = score ÷ 105 × 100. Use a separate answer booklet for working.')
        para(exam,'Suggested administration: two supervised sessions, 45 minutes for Section A and 135 minutes for Section B, or a school-approved equivalent. The requested 45 structured subparts make this a long assessment. Confirm timing and calculator policy before printing.')
        qs=[];essays=[]
        for i in range(15):
            week=i%maxweek;skills=WEEKS[g][week];chosen=skills[(i//maxweek)%len(skills)]
            t=math_mcq((chosen,),g,g*1000000+maxweek*10000+i*71,1)[0];t['source_week']=week+1;qs.append(t)
            parts=math_essays(skills,g,g*2000000+maxweek*10000+i*89,1)[0]
            for t in parts:t['source_week']=week+1
            essays.append(parts)
        put_mcq(exam,qs);put_essays(exam,essays,16);save(exam,folder/f'B{g}_Mathematics_{label}_Test.docx')
        k.add_page_break();k.add_heading(label.replace('_',' ')+' — Test answers',1)
        para(k,'Assessment blueprint, MCQ/essay → week: '+'; '.join(f'{i+1}/{i+16} → {i%maxweek+1}' for i in range(15)))
        mcq_key(k,qs);essay_key(k,essays,16)
        records.append(dict(grade=g,subject='Mathematics',test=label,mcq=qs,essays=essays))
    save(k,folder/f'B{g}_Mathematics_Teacher_Answers.docx')

def build_starter(g,subject,m):
    folder=OUT/f'B{g}'/subject
    d=newdoc(f'Basic {g} {subject} — Week 1',m['title']+' • Term 1 • Week 1 ONLY',True)
    para(d,'SCOPE: WEEK 1 ONLY. Later weeks and both tests for this subject are NOT included in this partial batch. Emergency teacher-review draft — rebuilt edition.')
    para(d,f'Sequence source: BASIC {g} SCHEME OF LEARNING.pdf, p{m["page"]}; indicator reference {m["codes"]}. Topic descriptions informed by the local curriculum data. Explanations and activities are newly prepared, including supporting concepts.')
    d.add_heading('Learning outcomes',1)
    for t in [f'Explain the key ideas in {m["title"].lower()} using accurate vocabulary.','Identify examples and justify classification or tool choices using evidence.','Apply the ideas to familiar school or home situations and correct a common misconception.']:bullet(d,t)
    d.add_heading('Teacher lesson notes',1)
    para(d,'One consolidated weekly unit. Suggested organisation: three 40-minute teaching sessions plus guided practice and assessment in the remaining timetable periods. Spread the questions over the week.')
    para(d,'Prior knowledge: elicit a familiar example and ask for a reason. Record one correct idea and one misconception to revisit at the end of the week.')
    for h,t in [('Session 1 — Observe','5 minutes: recall examples; 10: demonstrate or show pictures; 15: pairs describe features; 10: introduce the first five key terms and check explanations.'),('Session 2 — Investigate','5 minutes: retrieve vocabulary; 10: model an explanation; 15: perform the specific activity below; 10: compare observations and introduce the remaining terms.'),('Session 3 — Apply','5 minutes: correct a misconception; 10: explain map branches; 15: use application questions; 10: collect an exit response and identify who needs reteaching.')]:d.add_heading(h,2);para(d,t)
    d.add_heading('Topic-specific activity',2);para(d,m['activity'])
    d.add_heading('Safety and resources',2);para(d,m['safety']);para(d,'Resources: the examples or pictures specified in the activity, board, pencils, exercise books and the supplied visual map. If equipment is unavailable, use labelled drawings and role-play rather than claiming a practical was performed.')
    d.add_heading('Support and extension',2);para(d,'Support: offer picture labels, read technical vocabulary aloud and accept an oral explanation before writing. Extension: compare related concepts, propose a counterexample and justify a new application. Use wrong MCQ choices to diagnose misconceptions, not merely score letters.')
    d.add_heading('Study notes',1)
    for c in m['cards']:d.add_heading(c['answer'],2);para(d,c['fact']+'.');para(d,c['reason'])
    para(d,'Revision routine: cover each definition, explain the idea, supply a new example and compare it with a related term. Reconstruct the map from memory before checking it.')
    path=folder/'Mind-maps'/f'B{g}_{subject}_Week_01.png'
    mindmap(path,m['title'],g,subject,1,[(c['answer'],c['fact']+'.') for c in m['cards']],f'Sequence: BASIC {g} scheme p{m["page"]}; {m["codes"]}. Newly prepared concept map; teacher-review draft.')
    d.add_page_break();d.add_heading('Week 1 — Visual mind-map',1);d.add_picture(str(path),width=Inches(6.9));para(d,'Explain each branch and supply another example. A separate high-resolution PNG is included.')
    qs=[]
    for j,field in enumerate(['q1','q2']):
        for i,c in enumerate(m['cards']):qs.append(shuffled(dict(q=c[field],answer=c['answer'],wrong=c['wrong'],method=c['reason']),g*1000+len(subject)*100+j*10+i))
    d.add_page_break();d.add_heading('Week 1 — Practice',1);put_mcq(d,qs)
    d.add_heading('Section B — 10 structured essay questions',2)
    para(d,'Answer every part (a), (b), (c) in your exercise book. Suggested marks: (a) 1, (b) 1, (c) 2, giving 4 marks per question. Explain relationships and use examples where requested.')
    essays=[]
    for i,c in enumerate(m['cards'],1):
        questions=[f'Explain the term “{c["answer"]}”.',c['q2'],c['extend']]
        d.add_heading(f'Question {i}',3)
        for label,q in zip('abc',questions):para(d,f'({label}) {q}')
        essays.append([dict(q=q,answer=ans) for q,ans in zip(questions,[c['fact'],c['answer'],c['reason']])])
    save(d,folder/f'B{g}_{subject}_Week_01_Lesson_and_Practice.docx')
    k=newdoc(f'Basic {g} {subject} — Week 1 Teacher answers','Keep separate from pupil copy • Week 1 ONLY')
    mcq_key(k,qs);k.add_heading('Structured essay marking guide',1)
    para(k,'Part (a): 1 mark for an accurate definition. Part (b): 1 mark for correct identification. Part (c): 2 marks for two relevant points, or an accurate relationship with an appropriate example. Accept equivalent wording and other correct examples. Essay total: 40 marks.')
    for i,c in enumerate(m['cards'],1):k.add_heading(f'Question {i}',2);para(k,'(a) '+c['fact']+'.');para(k,'(b) '+c['answer']+'.');para(k,'(c) Indicative points: '+c['reason'])
    save(k,folder/f'B{g}_{subject}_Week_01_Teacher_Answers.docx')
    records.append(dict(grade=g,subject=subject,week=1,mcq=qs,essays=essays,map=str(path.relative_to(OUT))))

def guide():
    d=newdoc('START HERE — Delivery status','Rebuilt emergency Term 1 pack • PARTIAL DELIVERY • 22 September 2026')
    d.add_heading('What is included',1)
    para(d,'Mathematics B4–B8: Weeks 1–12, with weekly lesson notes, study notes, one visual mind-map per weekly unit, 20 four-option MCQs and 10 three-part structured questions per week. Each grade has a mid-term and end-of-term test, each with 15 MCQs and 15 three-part structured questions.')
    para(d,'Science and Computing B4–B8: WEEK 1 ONLY, each with lesson notes, study notes, one visual mind-map, 20 MCQs and 10 three-part structured questions. These are starter units, not full-term packs.')
    d.add_heading('What remains outstanding — NOT included',1)
    for t in ['Science: Weeks 2–12 for B4–B7, and Weeks 2–13 for B8. The B8 scheme contains new teaching topics in Week 13.','Computing: later teaching weeks for all five grades. B8 Week 12 is blank in the source and needs an agreed plan.','All Science and Computing mid-term and end-of-term tests.','Separate full daily lesson plans: these documents contain consolidated weekly units with suggested sessions.']:bullet(d,t)
    para(d,'No background production of the outstanding subjects is implied. The scope above is the complete scope of this archive.')
    d.add_heading('Files and use',1)
    para(d,'There are 41 editable Word files including this guide, and 70 standalone PNG mind-maps. Each Maths folder has a weekly book, two tests, a teacher key and 12 maps. Each Science and Computing folder has a Week 1 book, a teacher key and one map.')
    para(d,'Totals: 70 weekly units with 1,400 MCQs and 700 three-part questions; 10 Maths tests with 150 MCQs and 150 three-part questions. Every MCQ has four alternatives. Keep Teacher_Answers files away from pupil copies.')
    for t in ['Edit school details, dates, timetable, learner accommodations and assessment policy before printing.','Print weekly sections selectively; pupils use separate exercise books or answer booklets.','Maps are embedded in Word and supplied as high-resolution PNGs for projection or larger printing. They are connected concept maps, not reproductions of the plant-cell illustration.','Maths tests have 105 raw marks: 15 MCQ marks plus 90 structured-response marks. Convert to a percentage by multiplying score/105 by 100. The requested number of subparts requires extended or split administration.']:bullet(d,t)
    d.add_heading('Sources and editorial limitations',1)
    para(d,'Primary sequence: first-term-library/BASIC 4, 5, 6, 7 and 8 SCHEME OF LEARNING.pdf, Alpha Examinations annual schemes for 2026/2027. Supporting descriptions: data/curriculum/{subject}_B{grade}_curriculum_db_clean.json. First-term table pages: Maths 20/22/19/40/48; Science 24/26/23/55/62; Computing 54/54/51/15/21, in B4–B8 order.')
    para(d,'Source tables contain inconsistent or unavailable indicator suffixes and content-standard columns. B4 Maths Week 7 and B6 Maths Weeks 5 and 12 are examples. The pack follows interpreted topic intent rather than claiming exact code certification. B6 Week 12 uses rates/direct proportion as an editorial interpretation; B5 Week 5 includes LCM as support for HCF and factor–multiple relationships. Verify these choices before formal assessment.')
    para(d,'Maths uses topic-specific task templates with varied values and worked answers. Some structures intentionally repeat for fluency; not every item is pedagogically unique. Notes share explanations across grades where the concept is the same. Mid-term assumes Weeks 1–6 and end-of-term Weeks 1–12. Supporting starter concepts extend the central Week 1 topic where useful.')
    d.add_heading('Rebuilt edition and checks',1)
    para(d,'This archive was regenerated after the earlier download failure. It preserves the advertised partial scope, but is a rebuilt edition, not a claim to be byte-for-byte identical to the unavailable earlier archive.')
    para(d,'Checks cover DOCX integrity, document/question counts, four distinct alternatives, a single keyed answer, embedded maps and independent recalculation of a subset of numerical answers. See QUALITY_CHECKS.json for actual results. Word pagination has not been rendered in Microsoft Word here. These checks do not replace independent teacher review of difficulty, full curriculum coverage, language or accessibility.')
    save(d,OUT/'START_HERE_Delivery_Status.docx')
    (OUT/'README.txt').write_text('PARTIAL DELIVERY — REBUILT EDITION\n\nOpen START_HERE_Delivery_Status.docx first.\n\nIncluded: Maths B4-B8, Weeks 1-12 and both tests. Science and Computing B4-B8, WEEK 1 ONLY.\n\nOutstanding: later Science/Computing weeks and all their tests.\n\nAll included weekly units have a visual map, 20 MCQs, 10 three-part questions and separate teacher answers. Teacher-review drafts.\n',encoding='utf-8')

if __name__=='__main__':
    for g in range(4,9):build_math(g);print(f'Maths B{g} completed',flush=True)
    for (g,s),m in DATA.items():build_starter(g,s,m);print(f'{s} B{g} Week 1 completed',flush=True)
    guide();(ROOT/'source'/'assessment_records.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
    print('Generation finished',flush=True)

from pathlib import Path
import json,zipfile,re,math,hashlib
from fractions import Fraction
from collections import Counter
from docx import Document
from PIL import Image
ROOT=Path(__file__).resolve().parent.parent;OUT=ROOT/'files';DOWNLOADS=ROOT/'downloads';DOWNLOADS.mkdir(exist_ok=True)
records=json.loads((ROOT/'source/assessment_records.json').read_text())
weekly=[x for x in records if 'week' in x];tests=[x for x in records if 'test' in x]
assert len(weekly)==70 and len(tests)==10
assert len(list(OUT.rglob('*.docx')))==41 and len(list(OUT.rglob('*.png')))==70
mcq_count=essay_count=checked=0;balance=Counter()
for rec in records:
    expected=(20,10) if 'week' in rec else (15,15)
    assert (len(rec['mcq']),len(rec['essays']))==expected
    assert all(len(e)==3 for e in rec['essays'])
    if 'map' in rec:assert (OUT/rec['map']).exists()
    for q in rec['mcq']:
        assert len(q['options'])==4 and len(set(q['options']))==4
        assert q['options'].count(q['answer'])==1
        assert q['options']['ABCD'.index(q['letter'])]==q['answer']
        balance[q['letter']]+=1
    mcq_count+=len(rec['mcq']);essay_count+=len(rec['essays'])
    if rec['subject']!='Mathematics':continue
    for t in rec['mcq']+[t for e in rec['essays'] for t in e]:
        q=t['q'];ans=t['answer'];expected=None
        m=re.fullmatch(r'Calculate ([\d./]+) ([+−×÷]) ([\d./]+)\.(?: Give the answer in simplest form\.)?',q)
        if m:
            a=Fraction(m[1]);b=Fraction(m[3]);op=m[2]
            expected={'+':lambda:a+b,'−':lambda:a-b,'×':lambda:a*b,'÷':lambda:a/b}[op]()
        m=re.fullmatch(r'Find the (HCF|LCM) of (\d+) and (\d+)(?: using prime factors)?\.',q)
        if m:expected=Fraction(math.gcd(int(m[2]),int(m[3])) if m[1]=='HCF' else math.lcm(int(m[2]),int(m[3])))
        m=re.fullmatch(r'Evaluate (\d+)\^(\d+)\.',q)
        if m:expected=Fraction(int(m[1])**int(m[2]))
        m=re.fullmatch(r'Find ([\d/]+) of (\d+)\.',q)
        if m:expected=Fraction(m[1])*int(m[2])
        m=re.fullmatch(r'Find (\d+)% of (\d+)\.',q)
        if m:expected=Fraction(int(m[1])*int(m[2]),100)
        m=re.fullmatch(r'Find the non-negative square root of (\d+)\.',q)
        if m:expected=Fraction(math.isqrt(int(m[1])))
        if expected is not None:assert Fraction(ans)==expected,(q,ans,expected);checked+=1
embedded=0
for p in OUT.rglob('*.docx'):
    with zipfile.ZipFile(p) as z:
        assert z.testzip() is None
        embedded+=len([n for n in z.namelist() if n.startswith('word/media/')])
    d=Document(p);text='\n'.join(x.text for x in d.paragraphs)
    assert text.strip() and 'TODO' not in text
    if 'Weekly_Lessons' in p.name:
        assert len(d.inline_shapes)==12
        assert sum(x.text=='Section A — Multiple-choice questions' for x in d.paragraphs)==12
        assert sum(x.text=='Section B — Structured essay questions' for x in d.paragraphs)==12
        assert 'MCQ answers and explanations' not in text
    if 'Lesson_and_Practice' in p.name:assert len(d.inline_shapes)==1
    if p.name.endswith('_Test.docx'):
        assert len([x for x in d.paragraphs if re.match(r'^\d+\. ',x.text)])==30
        assert len([x for x in d.paragraphs if re.match(r'^\([abc]\) ',x.text)])==45
for p in OUT.rglob('*.png'):
    im=Image.open(p);assert im.size==(2200,1800);im.verify()
assert embedded==70
report={'edition':'Rebuilt after download failure; not byte-identical to previous archive','scope':'PARTIAL: Maths B4–B8 Weeks 1–12 and both tests; Science/Computing B4–B8 Week 1 only','word_documents':41,'weekly_units':70,'mindmaps':70,'embedded_mindmaps':embedded,'mathematics_tests':10,'weekly_mcqs':1400,'weekly_three_part_questions':700,'test_mcqs':150,'test_three_part_questions':150,'total_mcqs':mcq_count,'total_three_part_questions':essay_count,'independently_recomputed_numeric_answers':checked,'answer_position_distribution':dict(balance),'checks':'DOCX integrity; counts; four distinct alternatives; one keyed answer; embedded images; test numbering; partial independent numeric recomputation; image integrity.','limits':'Not independently teacher-certified; no full Word pagination render; later Science/Computing weeks and all their tests outstanding.'}
(OUT/'QUALITY_CHECKS.json').write_text(json.dumps(report,indent=2,ensure_ascii=False),encoding='utf-8')
(OUT/'FILE_MANIFEST.txt').write_text('\n'.join(f'{p.relative_to(OUT)} | {p.stat().st_size:,} bytes' for p in sorted(OUT.rglob('*')) if p.is_file() and p.name!='FILE_MANIFEST.txt')+'\n',encoding='utf-8')
zip_path=DOWNLOADS/'Beacon_Term1_Batch1_Rebuilt_PARTIAL.zip'
with zipfile.ZipFile(zip_path,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
    for p in sorted(OUT.rglob('*')):
        if p.is_file():z.write(p,arcname=str(p.relative_to(OUT)))
with zipfile.ZipFile(zip_path) as z:assert z.testzip() is None
info={'filename':zip_path.name,'bytes':zip_path.stat().st_size,'sha256':hashlib.sha256(zip_path.read_bytes()).hexdigest()}
(DOWNLOADS/'archive-info.json').write_text(json.dumps(info,indent=2))
print(json.dumps(report,indent=2));print(json.dumps(info,indent=2))

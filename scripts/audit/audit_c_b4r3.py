#!/usr/bin/env python3
"""Audit C (B4 remaining-3) - Verify Creative Arts, History, RME B4 lesson JSONs + DOCX."""
import json, re
from collections import Counter
import docx
from docx.shared import RGBColor
from pathlib import Path  # noqa: E402  (legacy script)

ROOT = str(Path(__file__).resolve().parents[2])
WEEKDAYS = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday']
CODE_RE = re.compile(r'^B4.\d+\.\d+\.\d+\.\d+$')
REQ = ['strand_name', 'sub_strand', 'cs_code', 'cs_desc', 'ind_code', 'ind_desc',
       'perf_indicator', 'competencies', 'resources', 'starter', 'main', 'plenary', 'assessment']

SUBJECTS = {
    'creative_arts_b4': ('curriculum_db/creative-arts_B4_curriculum_db_clean.json', 'Basic4_Creative_Arts_Lesson_Plans_Full_Year.docx'),
    'history_b4':       ('curriculum_db/history_B4_curriculum_db_clean.json',      'Basic4_History_Lesson_Plans_Full_Year.docx'),
    'rme_b4':           ('curriculum_db/rme_B4_curriculum_db_clean.json',          'Basic4_RME_Lesson_Plans_Full_Year.docx'),
}

def audit_json(stem, dbp):
    issues = []
    lessons = json.load(open(f'{ROOT}/{stem}_lessons_enriched.json'))
    db = json.load(open(f'{ROOT}/{dbp}'))
    if len(lessons) != 180:
        issues.append(f'expected 180, got {len(lessons)}')
    if sorted(l['lesson_num'] for l in lessons) != list(range(1, 181)):
        issues.append('lesson_num not 1..180')
    slots = Counter((l['term'], l['week'], l['day']) for l in lessons)
    expected = {(t, w, d) for t in (1, 2, 3) for w in range(1, 13) for d in WEEKDAYS}
    if [s for s, c in slots.items() if c > 1]: issues.append('duplicate slots')
    if expected - set(slots): issues.append(f'{len(expected - set(slots))} empty slots')
    if set(slots) - expected: issues.append('invalid slots')
    bad_code = [l['lesson_num'] for l in lessons if not CODE_RE.match(l['ind_code'])]
    not_in_db = [(l['lesson_num'], l['ind_code']) for l in lessons if l['ind_code'] not in db]
    missing = [(l['lesson_num'], f) for l in lessons for f in REQ
               if l.get(f) is None or len(l.get(f)) == 0]
    placeholder = [l['lesson_num'] for l in lessons if 'Learning Indicator' in str(l['ind_desc'])]
    # cs_desc must not be placeholder either
    ph_cs = [l['lesson_num'] for l in lessons if 'Content Standard' in str(l['cs_desc'])]
    if bad_code: issues.append(f'malformed codes: {bad_code[:5]}')
    if not_in_db: issues.append(f'codes not in DB: {not_in_db[:5]}')
    if missing: issues.append(f'empty fields: {missing[:5]}')
    if placeholder: issues.append(f'placeholder ind_desc: {placeholder[:5]}')
    if ph_cs: issues.append(f'placeholder cs_desc: {ph_cs[:5]}')
    return {'subject': stem, 'status': 'PASS' if not issues else 'FAIL', 'issues': issues}

def audit_docx(stem, docx_name, lessons):
    issues = []
    d = docx.Document(f'{ROOT}/{docx_name}')
    j = '\n'.join(p.text for p in d.paragraphs)
    if len(re.findall(r'LESSON PLAN #\d+', j)) != 180:
        issues.append('headers != 180')
    heads = [(int(a), int(b), c) for a, b, c in
             re.findall(r'LESSON PLAN #(\d+)\s*\xb7\s*WEEK\s*(\d+)\s*\xb7\s*([A-Z]+)', j)]
    want = [(l['lesson_num'], l['week'], l['day'].upper()) for l in lessons]
    if heads != want:
        issues.append(f'header seq mismatch (parsed {len(heads)})')
    for label, pat in [('PHASE 1', r'PHASE 1\s*—\s*STARTER'), ('PHASE 2', r'PHASE 2\s*—\s*MAIN'),
                       ('PHASE 3', r'PHASE 3\s*—\s*PLENARY'), ('Reflection', r"Teacher's Reflection"),
                       ('Sign-off', r"Teacher's Signature")]:
        c = len(re.findall(pat, j))
        if c != 180: issues.append(f'{label}: {c} != 180')
    if 'Basic 4' not in j: issues.append('title does not say Basic 4')
    fonts, navy = set(), False
    for p in d.paragraphs[:400]:
        for r in p.runs:
            if r.font.name: fonts.add(r.font.name)
            try:
                if r.font.color and r.font.color.rgb == RGBColor(0, 32, 96): navy = True
            except Exception: pass
    if fonts - {'Calibri'}: issues.append(f'fonts: {sorted(fonts)}')
    if not navy: issues.append('navy color missing')
    return {'subject': stem, 'docx': docx_name, 'status': 'PASS' if not issues else 'FAIL', 'issues': issues}

rows = [audit_json(s, d) for s, (d, _) in SUBJECTS.items()]
for s, (_, docx_name) in SUBJECTS.items():
    rows.append(audit_docx(s, docx_name, json.load(open(f'{ROOT}/{s}_lessons_enriched.json'))))
ok = True
for r in rows:
    print(f"[{r['status']}] {r['subject']:18s} {r.get('docx','(json)')}")
    for i in r['issues']: print('   !', i); ok = False
print('\nRESULT:', 'ALL PASS' if ok else 'FAILURES PRESENT')

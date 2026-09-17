#!/usr/bin/env python3
"""Generate Records of Work from the enriched lesson library.

Where a Scheme of Learning is the *plan* (submitted at the start of term), the
Record of Work is the *log* — filled in as teaching happens and signed by the
headteacher. It is the third artifact in the chain, and the one that closes the
loop: plan → teach → record.

Because the lesson library already knows what is taught on every day of the
year, the Record of Work can be pre-printed with the curriculum references and
activity summary for every period, leaving the teacher only the parts that must
be handwritten: date, evaluation and remarks.

Sources
-------
    {subject}[_b{grade}]_lessons_enriched.json    (73 files, 180 lessons each)

Outputs
-------
    dist/records/docx/Record_of_Work_{Subject}_Basic{N}.docx
        One ledger per subject-grade (73 total), landscape, three terms.

    dist/records/docx/terms/…   (with --per-term)
        One document per subject-grade-term (219 total).

    dist/records/json/{subjectId}_{grade}_record.json
        Day-level rows, for a future in-app Record of Work module.

    dist/records/QA_REPORT.md

Requires: python-docx  (pip install python-docx)

Usage
-----
    python3 tools/generate_records_of_work.py
    python3 tools/generate_records_of_work.py --grade B4
    python3 tools/generate_records_of_work.py --per-term
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))

from generate_schemes import (  # noqa: E402
    DAY_ORDER,
    HAS_DOCX,
    SUBJECTS,
    add_rule,
    cell_text,
    discover,
    field,
    set_widths,
    shade,
    styled,
)

if not HAS_DOCX:
    sys.exit("python-docx is required:  pip install python-docx")

from docx import Document  # noqa: E402
from docx.enum.section import WD_ORIENT  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Inches, Pt  # noqa: E402

OUT = ROOT / "dist" / "records"
JSON_OUT = OUT / "json"
DOCX_OUT = OUT / "docx"

NAVY = (0, 32, 96)
WEEK_FILL = "E2E8F0"

# Landscape Letter (11" x 8.5") with 0.5" side margins -> 10.0" usable.
# Keep this sum at or under USABLE_WIDTH — an over-wide table silently runs off
# the printed page, which is easy to miss until someone prints it.
PAGE_SHORT_SIDE = 8.5
SIDE_MARGIN = 0.5
USABLE_WIDTH = 11.0 - (2 * SIDE_MARGIN)  # 10.0

COL_WIDTHS = [0.55, 0.62, 0.29, 0.35, 1.27, 1.23, 1.41, 1.70, 0.95, 1.63]

assert len(COL_WIDTHS) == 10
assert abs(sum(COL_WIDTHS) - USABLE_WIDTH) < 0.01, (
    f"Record of Work columns sum to {sum(COL_WIDTHS):.2f}"
    f" but only {USABLE_WIDTH:.2f} in is usable"
)
COL_HEADS = [
    "DATE", "DAY", "WK", "LES", "STRAND / SUB-STRAND",
    "CONTENT STANDARD & INDICATOR", "LEARNING OUTCOME(S)",
    "ACTIVITIES UNDERTAKEN", "T / L RESOURCES", "EVALUATION / REMARKS",
]

BANDS = {
    **{f"B{i}": "Lower Primary" for i in (1, 2, 3)},
    **{f"B{i}": "Upper Primary" for i in (4, 5, 6)},
    "B7": "JHS 1", "B8": "JHS 2", "B9": "JHS 3",
}


def clean(v):
    if not isinstance(v, str):
        return ""
    v = re.sub(r"\s*===\s*PAGE \d+\s*===.*$", "", v, flags=re.S)
    return re.sub(r"\s+", " ", v).strip()


def trunc(text, n):
    text = clean(text)
    return text if len(text) <= n else text[: n - 1].rstrip() + "…"


def outcome(lesson):
    """Strip the boilerplate prefix from the performance indicator."""
    p = clean(lesson.get("perf_indicator"))
    p = re.sub(r"^By the end of the lesson,?\s*learners?\s*will be able to:?\s*",
               "", p, flags=re.I)
    return trunc(p, 150)


def activities(lesson):
    """A brief, printable summary of what the period covers."""
    main = lesson.get("main") or []
    plen = lesson.get("plenary") or []
    out = []
    if main:
        out.append(trunc(main[0], 130))
    if len(main) > 1:
        out.append(trunc(main[1], 130))
    if plen:
        out.append(trunc(plen[0], 95))
    return out


def build_rows(lessons, term):
    """Week separator rows + one row per teaching day, for one term."""
    by_week = {}
    for l in lessons:
        if l.get("term") == term:
            by_week.setdefault(l.get("week"), []).append(l)

    rows = []
    for week in sorted(by_week):
        rows.append({"kind": "week", "week": week})
        for l in sorted(by_week[week],
                        key=lambda x: DAY_ORDER.index(x["day"])
                        if x.get("day") in DAY_ORDER else 99):
            rows.append({
                "kind": "day",
                "week": week,
                "day": l.get("day", ""),
                "lessonNum": l.get("lesson_num"),
                "indicatorId": f"{SUBJECTS_KEY}_{l.get('ind_code', '')}",
                "strand": clean(l.get("strand_name")),
                "subStrand": clean(l.get("sub_strand")),
                "contentStandardCode": l.get("cs_code", ""),
                "indicatorCode": l.get("ind_code", ""),
                "outcome": outcome(l),
                "activities": activities(l),
                "resources": trunc(l.get("resources"), 110),
                "isRevision": bool(l.get("is_revision")),
            })
    return rows


SUBJECTS_KEY = ""  # set per-run so indicatorId can be namespaced


# ------------------------------------------------------------------ document

def new_doc():
    doc = Document()
    s = doc.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Inches(11), Inches(8.5)
    s.left_margin = s.right_margin = Inches(0.5)
    s.top_margin = s.bottom_margin = Inches(0.6)
    return doc


def table_for(doc, rows):
    table = doc.add_table(rows=1, cols=len(COL_HEADS))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr = table.rows[0]
    for i, head in enumerate(COL_HEADS):
        cell_text(hdr.cells[i], head, size=6.5, bold=True)
        shade(hdr.cells[i], WEEK_FILL)

    for row in rows:
        cells = table.add_row().cells
        if row["kind"] == "week":
            cell_text(cells[0], f"WEEK {row['week']}", size=7, bold=True)
            cell_text(cells[1], f"Week Ending: ______________________",
                      size=7, bold=True)
            cells[1].merge(cells[len(COL_HEADS) - 1])
            shade(cells[1], WEEK_FILL)
            for c in cells[2:]:
                pass
        else:
            cell_text(cells[0], "", size=7)                      # date (blank)
            cell_text(cells[1], row["day"], size=7)
            cell_text(cells[2], str(row["week"]), size=6.5)
            cell_text(cells[3], str(row["lessonNum"] or ""), size=6.5)
            cell_text(cells[4],
                      "\n".join(x for x in [row["strand"], row["subStrand"]] if x),
                      size=6.5)
            cell_text(cells[5],
                      "\n".join([row["contentStandardCode"], row["indicatorCode"],
                                 trunc(row.get("indicatorText", ""), 95)]),
                      size=6.5)
            cell_text(cells[6], row["outcome"], size=6.5)
            cell_text(cells[7], "\n".join(row["activities"]), size=6.5)
            cell_text(cells[8], row["resources"], size=6.5)
            cell_text(cells[9], "", size=6.5)                    # evaluation (blank)

    set_widths(table, COL_WIDTHS)
    return table


def cover(doc, subject_name, grade, school=None, teacher=None,
          class_name=None, term=None, year=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    styled(p, "RECORD OF WORK", size=22, bold=True, color=__navy())

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    styled(p, f"{subject_name}  ·  Basic {grade[1:]}  ·  {BANDS.get(grade, '')}",
           size=12, bold=True)

    for label in [
        f"{field('School', school, 38)}      "
        f"{field('Class', class_name, 12)}      {field('Academic Year', year, 10)}",
        f"{field('Teacher', teacher, 36)}      "
        f"Subject: {subject_name}      {field('Term', term, 6)}",
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        styled(p, label, size=9, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    styled(p,
           "Pre-printed from the Beacon lesson library: every teaching day of the "
           "year is listed with its curriculum reference and activity summary. "
           "Complete the DATE, EVALUATION / REMARKS column as you teach, and have "
           "the headteacher sign at the end of each week.",
           size=8, italic=True)
    add_rule(p, "002060", "8")


def __navy():
    from docx.shared import RGBColor
    return RGBColor(*NAVY)


def signatures(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    styled(p,
           "Class Teacher's Signature: ____________________      "
           "Date: ____________            "
           "Headteacher's Signature: ____________________      Date: ____________",
           size=9, bold=True)


def build_docx(subject_name, grade, terms_rows, path, **branding):
    doc = new_doc()
    cover(doc, subject_name, grade, **branding)
    for term in sorted(terms_rows, key=lambda t: int(t)):
        doc.add_page_break()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        styled(p, f"TERM {term}  ·  {subject_name.upper()}  ·  BASIC {grade[1:]}",
               size=11, bold=True, color=__navy())
        table_for(doc, terms_rows[term])
        signatures(doc)
    doc.save(path)
    return path


# ----------------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--grade")
    ap.add_argument("--subject")
    ap.add_argument("--per-term", action="store_true")
    ap.add_argument("--out")
    # Branding — pre-print the cover for a specific school (optional).
    ap.add_argument("--school", help="pre-print the school name on the cover")
    ap.add_argument("--teacher", help="pre-print the teacher's name")
    ap.add_argument("--class-name", dest="class_name",
                    help="pre-print the class, e.g. 'Basic 4'")
    ap.add_argument("--term", help="pre-print the term")
    ap.add_argument("--year", help="pre-print the academic year")
    args = ap.parse_args()

    branding = {k: v for k, v in dict(
        school=args.school, teacher=args.teacher, class_name=args.class_name,
        term=args.term, year=args.year).items() if v}

    global OUT, JSON_OUT, DOCX_OUT
    if args.out:
        OUT = Path(args.out); JSON_OUT = OUT / "json"; DOCX_OUT = OUT / "docx"
    for d in (JSON_OUT, DOCX_OUT):
        d.mkdir(parents=True, exist_ok=True)
    if args.per_term:
        (DOCX_OUT / "terms").mkdir(parents=True, exist_ok=True)

    entries = discover()
    if args.grade:
        entries = [e for e in entries if e[1] == args.grade]
    if args.subject:
        entries = [e for e in entries if e[0] == args.subject]
    if not entries:
        sys.exit("No matching lesson files.")

    print(f"Building Records of Work for {len(entries)} subject-grade file(s)...\n")

    qa = []
    for subj, grade, path in entries:
        global SUBJECTS_KEY
        subject_id, subject_name = SUBJECTS[subj]
        SUBJECTS_KEY = subject_id

        lessons = json.loads(path.read_text(encoding="utf-8"))
        terms_rows = {t: build_rows(lessons, t)
                      for t in sorted({l["term"] for l in lessons})}
        # Attach the indicator description by code (rows only carry codes).
        desc_by_code = {l.get("ind_code"): clean(l.get("ind_desc"))
                        for l in lessons}
        for rows in terms_rows.values():
            for r in rows:
                if r["kind"] == "day":
                    r["indicatorText"] = desc_by_code.get(r["indicatorCode"], "")

        safe = subject_name.replace(" ", "_").replace("&", "and")
        docx_path = DOCX_OUT / f"Record_of_Work_{safe}_Basic{grade[1:]}.docx"
        build_docx(subject_name, grade, terms_rows, docx_path, **branding)

        if args.per_term:
            for term, rows in terms_rows.items():
                d = new_doc()
                p = d.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                styled(p, "RECORD OF WORK", size=16, bold=True, color=__navy())
                p = d.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(6)
                styled(p, f"{subject_name}  ·  Basic {grade[1:]}  ·  Term {term}",
                       size=11, bold=True)
                p = d.add_paragraph()
                p.paragraph_format.space_after = Pt(4)
                styled(p, f"{field('School', args.school, 30)}      "
                          f"{field('Class', args.class_name, 10)}      "
                          f"{field('Teacher', args.teacher, 22)}      "
                          f"{field('Year', args.year, 10)}", size=9, bold=True)
                table_for(d, rows)
                signatures(d)
                d.save(DOCX_OUT / "terms" /
                       f"Record_of_Work_{safe}_Basic{grade[1:]}_Term{term}.docx")

        (JSON_OUT / f"{subject_id}_{grade}_record.json").write_text(json.dumps({
            "subjectId": subject_id,
            "subject": subject_name,
            "grade": grade,
            "source": path.name,
            "terms": {str(t): rows for t, rows in sorted(terms_rows.items())},
        }, indent=2, ensure_ascii=False), encoding="utf-8")

        days = sum(1 for rows in terms_rows.values()
                   for r in rows if r["kind"] == "day")
        weeks = sum(1 for rows in terms_rows.values()
                    for r in rows if r["kind"] == "week")
        no_outcome = sum(1 for rows in terms_rows.values() for r in rows
                         if r["kind"] == "day" and not r["outcome"])
        qa.append({
            "grade": grade, "subject": subject_name,
            "days": days, "weeks": weeks,
            "lessons_in": len(lessons),
            "no_outcome": no_outcome,
            "docx_kb": round(docx_path.stat().st_size / 1024),
            "status": "OK" if days == len(lessons) else "CHECK",
        })
        print(f"  {grade} {subject_name:32} {days:3} days · {weeks:2} weeks · "
              f"{qa[-1]['docx_kb']:4} KB  [{qa[-1]['status']}]")

    write_qa(qa)
    print(f"\nDone.\n  DOCX : {DOCX_OUT}\n  JSON : {JSON_OUT}\n  QA   : {OUT / 'QA_REPORT.md'}")


def write_qa(qa):
    bad = [r for r in qa if r["status"] != "OK"]
    lines = [
        "# Records of Work — QA Report",
        "",
        f"Generated {len(qa)} subject-grade records "
        f"({sum(r['days'] for r in qa):,} teaching days across "
        f"{sum(r['weeks'] for r in qa):,} weeks).",
        "",
    ]
    lines.append("> ✅ Every lesson in the library appears as a teaching-day row."
                 if not bad else
                 f"> ⚠️ **{len(bad)} record(s) need attention.**")
    lines += [
        "",
        "| Grade | Subject | Days | Weeks | Lessons in | Missing outcomes | DOCX | Status |",
        "|---|---|---|---|---|---|---|---|",
    ]
    for r in sorted(qa, key=lambda x: (int(x["grade"][1:]), x["subject"])):
        lines.append(f"| {r['grade']} | {r['subject']} | {r['days']} | {r['weeks']} | "
                     f"{r['lessons_in']} | {r['no_outcome']} | {r['docx_kb']} KB | "
                     f"{r['status']} |")
    lines += [
        "",
        "## Notes",
        "",
        "- **Days** must equal **Lessons in** (180) — one ledger row per teaching day.",
        "- DATE and EVALUATION / REMARKS are left blank deliberately: those are the",
        "  parts only the teacher can write.",
        "- Each week is preceded by a shaded WEEK n / Week Ending separator row, so",
        "  the headteacher can sign week by week.",
        "- Landscape Letter, 0.5\" margins, 6.5pt table text — sized to be printed",
        "  and bound as a ledger.",
    ]
    (OUT / "QA_REPORT.md").write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    main()
